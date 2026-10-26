"""在 Word 作业上落笔：在每道题的内容末尾追加一段红色分数（可编辑文本）。

用法：
    python annotate_docx.py <成绩json> <原docx> <输出docx>
"""

import json
import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph

sys.stdout.reconfigure(encoding="utf-8")

RED = RGBColor(0xC0, 0x00, 0x00)


def is_heading(p):
    name = (p.style.name or "").lower()
    return "heading" in name or "标题" in name


def find_paragraph(doc, anchor):
    """优先匹配标题开头的段落，其次全文包含。"""
    anchor = (anchor or "").strip()
    if not anchor:
        return None
    probe = anchor[:10]
    for p in doc.paragraphs:
        if p.text.strip().startswith(probe):
            return p
    for p in doc.paragraphs:
        if probe in p.text:
            return p
    return None


def section_end(doc, start):
    """题目内容末尾：从标题下一段起，到下一个标题或文末为止。"""
    paras = doc.paragraphs
    idx = next(i for i, p in enumerate(paras) if p._p is start._p)
    end = start
    for p in paras[idx + 1:]:
        if is_heading(p):
            break
        end = p
    return end


def add_paragraph_after(ref, text, size_pt):
    new_p = OxmlElement("w:p")
    ref._p.addnext(new_p)
    p = Paragraph(new_p, ref._parent)
    run = p.add_run(text)
    run.font.color.rgb = RED
    run.bold = True
    run.font.size = Pt(size_pt)
    return p


def main(grades_path, src, dst):
    with open(grades_path, encoding="utf-8") as f:
        data = json.load(f)

    doc = Document(src)
    missed = []
    for item in data.get("items", []):
        score = item.get("score")
        if score is None:
            continue
        anchor = item.get("anchor") or item.get("title") or ""
        head = find_paragraph(doc, anchor)
        if head is None:
            missed.append(anchor)
            continue

        base = head.runs[0].font.size.pt if head.runs and head.runs[0].font.size else 12
        label = f"{item.get('title', '得分')}：{score}"
        if item.get("max"):
            label += f"/{item['max']}"
        comment = (item.get("comment") or "").strip()
        if comment:
            label += f"（{comment}）"
        add_paragraph_after(section_end(doc, head), label, base)

    total = data.get("total")
    if total:
        # 总分写在文档开头，右对齐
        p = doc.paragraphs[0].insert_paragraph_before() if doc.paragraphs else doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        text = f"总分：{total['score']}" + (f"/{total['max']}" if total.get("max") else "")
        run = p.add_run(text)
        run.font.color.rgb = RED
        run.bold = True
        run.font.size = Pt(14)

    os.makedirs(os.path.dirname(os.path.abspath(dst)) or ".", exist_ok=True)
    doc.save(dst)

    print(f"已生成：{dst}")
    if missed:
        print(f"  以下评分点未在文中定位，未落笔：{', '.join(missed)}")


if __name__ == "__main__":
    if len(sys.argv) < 4:
        print(__doc__)
        sys.exit(1)
    main(sys.argv[1], sys.argv[2], sys.argv[3])
