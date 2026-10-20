"""汇总所有成绩 json，生成 Excel 统计表。

用法：
    python make_excel.py <grades目录> <输出xlsx>

表格列：文件名 | 各评分项得分 | 总分 | 满分 | 备注
总分由各分项相加得出，不采信模型给的总分（不一致时会提示）。
"""

import glob
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8")


def load(grades_dir):
    rows = []
    titles = []
    for path in sorted(glob.glob(os.path.join(grades_dir, "*.json"))):
        with open(path, encoding="utf-8") as f:
            data = json.load(f)
        items = data.get("items") or []
        for it in items:
            if it["title"] not in titles:
                titles.append(it["title"])
        rows.append(data)
    return rows, titles


def main(grades_dir, out):
    rows, titles = load(grades_dir)
    if not rows:
        raise SystemExit(f"{grades_dir} 下没有成绩 json")

    wb = Workbook()
    ws = wb.active
    ws.title = "批改统计"

    headers = ["文件名"] + titles + ["总分", "满分", "备注"]
    ws.append(headers)

    problems = []
    for data in rows:
        scores = {it["title"]: it["score"] for it in data.get("items", [])}
        total = sum(v for v in scores.values() if isinstance(v, (int, float)))
        declared = (data.get("total") or {}).get("score")
        if declared is not None and abs(declared - total) > 0.01:
            problems.append(f"{data['filename']}：模型报 {declared}，分项合计 {total}，已按分项计")
        full = (data.get("total") or {}).get("max") or sum(
            it.get("max", 0) for it in data.get("items", [])
        )
        ws.append(
            [data["filename"]]
            + [scores.get(t, "") for t in titles]
            + [total, full, data.get("notes", "")]
        )

    # 按总分降序，方便看全班分布
    body = list(ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True))
    body.sort(key=lambda r: r[len(titles) + 1], reverse=True)
    for i, row in enumerate(body, start=2):
        for j, v in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=v)

    thin = Side(style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="F2F2F2")
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col)
        c.font = Font(bold=True)
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=len(headers)):
        for c in row:
            c.border = border
            if c.column > 1:
                c.alignment = Alignment(horizontal="center")

    ws.column_dimensions["A"].width = max(28, *(len(str(r[0])) for r in body)) + 4
    for i in range(2, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = max(10, len(str(headers[i - 1])) * 2 + 4)
    ws.freeze_panes = "B2"

    os.makedirs(os.path.dirname(os.path.abspath(out)) or ".", exist_ok=True)
    wb.save(out)

    avg = sum(r[len(titles) + 1] for r in body) / len(body)
    print(f"已生成：{out}")
    print(f"  共 {len(body)} 份，平均分 {avg:.1f}")
    for p in problems:
        print(f"  ⚠ {p}")


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(1)
    main(sys.argv[1], sys.argv[2])
